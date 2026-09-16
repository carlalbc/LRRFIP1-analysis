"""
Create Excel tables for:
1. Main figure data: short-read and long-read isoform fractions per individual
2. Linked genes TPM: fibroblast (short-read + long-read) and optional iPSC
"""
import pandas as pd
import numpy as np
import os, glob
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

OUTDIR = "./data/salmon_quant_iPSC"
FIBRO_ROOT = "./data/salmon_quant_fibroblast"
IPSC_ROOT = OUTDIR
TX2GENE = os.path.join(OUTDIR, "tx2gene_gencode49.tsv")

ISO_ORDER = ["Isoform 1", "Isoform 2", "Isoform 3", "Isoform 4", "Isoform 5", "Isoform 6"]
ENST_COLS = {
    "ENST00000308482.14": "Isoform 1",
    "ENST00000863709.1":  "Isoform 2",
    "ENST00000392000.4":  "Isoform 3",
    "ENST00000244815.9":  "Isoform 4",
    "ENST00000289175.10": "Isoform 5",
    "ENST00000478958.5":  "Isoform 6",
}

INDIV_ORDER = ["II-8.5", "II-8", "III-13", "III-12"]
INDIV_GENO = {"III-13": "WT", "II-8.5": "WT", "III-12": "Affected", "II-8": "Affected"}
INDIV_GEN = {"II-8.5": "Gen. II", "II-8": "Gen. II", "III-13": "Gen. III", "III-12": "Gen. III"}
INDIV_RELATION = {"II-8.5": "Mother", "II-8": "Father", "III-13": "Son", "III-12": "Son"}

# Pedigree to sample mapping
def get_individual(s):
    if "33016" in s: return "III-13"
    elif "33017" in s: return "II-8.5"
    elif "33018" in s: return "III-12"
    elif "33019" in s: return "II-8"
    return s

# ══════════════════════════════════════════════════
# TABLE 1: Short-read isoform fractions
# ══════════════════════════════════════════════════
tpm = pd.read_csv(f"{FIBRO_ROOT}/lrrfip1_iso6_summary/LRRFIP1_isoforms_TPM.tsv", sep="\t")
tpm = tpm.set_index("sample")
tpm.columns = [ENST_COLS[c] for c in tpm.columns]
tpm["individual"] = tpm.index.map(lambda s: get_individual(s))
tpm["genotype"] = tpm["individual"].map(INDIV_GENO)
tpm["total"] = tpm[ISO_ORDER].sum(axis=1)
for iso in ISO_ORDER:
    tpm[f"{iso}_frac"] = tpm[iso] / tpm["total"]

sr_rows = []
for ind in INDIV_ORDER:
    sub = tpm[tpm["individual"] == ind]
    row = {
        "Individual": ind,
        "Generation": INDIV_GEN[ind],
        "Relation": INDIV_RELATION[ind],
        "Genotype": INDIV_GENO[ind],
        "n replicates": len(sub),
        "Total LRRFIP1 TPM (mean)": sub["total"].mean(),
    }
    for iso in ISO_ORDER:
        row[f"{iso} fraction (mean)"] = sub[f"{iso}_frac"].mean()
        row[f"{iso} fraction (SD)"] = sub[f"{iso}_frac"].std()
        row[f"{iso} TPM (mean)"] = sub[iso].mean()
    sr_rows.append(row)

sr_df = pd.DataFrame(sr_rows)

# ══════════════════════════════════════════════════
# TABLE 2: Long-read (IsoQuant FLNC) isoform data
# ══════════════════════════════════════════════════
lr_counts = {
    "Isoform 1": {"III-13": 0,   "II-8.5": 0,   "III-12": 0,   "II-8": 0},
    "Isoform 2": {"III-13": 111, "II-8.5": 112,  "III-12": 92,  "II-8": 165},
    "Isoform 3": {"III-13": 0,   "II-8.5": 3,    "III-12": 0,   "II-8": 0},
    "Isoform 4": {"III-13": 159, "II-8.5": 119,  "III-12": 87,  "II-8": 82},
    "Isoform 5": {"III-13": 346, "II-8.5": 213,  "III-12": 189, "II-8": 236},
    "Isoform 6": {"III-13": 29,  "II-8.5": 17,   "III-12": 31,  "II-8": 16},
    "Other":     {"III-13": 37,  "II-8.5": 44,   "III-12": 40,  "II-8": 37},
}
all_lr_isos = ISO_ORDER + ["Other"]

lr_rows = []
for ind in INDIV_ORDER:
    total = sum(lr_counts[iso][ind] for iso in all_lr_isos)
    row = {
        "Individual": ind,
        "Generation": INDIV_GEN[ind],
        "Relation": INDIV_RELATION[ind],
        "Genotype": INDIV_GENO[ind],
        "Total FLNC reads": total,
    }
    for iso in all_lr_isos:
        row[f"{iso} reads"] = lr_counts[iso][ind]
        row[f"{iso} fraction"] = lr_counts[iso][ind] / total if total > 0 else 0
    lr_rows.append(row)

lr_df = pd.DataFrame(lr_rows)

# ══════════════════════════════════════════════════
# TABLE 3: Linked genes TPM (fibroblasts, short-read + long-read)
# ══════════════════════════════════════════════════
tx2g = pd.read_csv(TX2GENE, sep="\t", header=None, names=["transcript_id", "gene_name"])
tx_map = dict(zip(tx2g["transcript_id"], tx2g["gene_name"]))

# Genomic order on chr2: COL6A3, MLPH, PRLH, RAB17, LRRFIP1, RBM44
LINKED_GENES = ["COL6A3", "MLPH", "PRLH", "RAB17", "LRRFIP1", "RBM44"]

def load_gene_tpm(quant_files, tx_map, label_func):
    records = []
    for qf in sorted(quant_files):
        sample = label_func(qf)
        df = pd.read_csv(qf, sep="\t")
        df["gene"] = df["Name"].map(tx_map)
        df = df.dropna(subset=["gene"])
        gene_tpm = df.groupby("gene")["TPM"].sum().reset_index()
        gene_tpm["sample"] = sample
        records.append(gene_tpm)
    return pd.concat(records, ignore_index=True)

# Fibroblast short-read gene-level TPM
fibro_files = glob.glob(os.path.join(FIBRO_ROOT, "*/quant.sf"))
fibro_data = load_gene_tpm(fibro_files, tx_map, lambda f: os.path.basename(os.path.dirname(f)))
fibro_data["individual"] = fibro_data["sample"].apply(get_individual)
fibro_data["genotype"] = fibro_data["individual"].map(INDIV_GENO)

# Per-individual means (mean of replicate means)
linked_rows_fibro = []
for gene in LINKED_GENES:
    row = {"Gene": gene}
    for ind in INDIV_ORDER:
        sub = fibro_data[(fibro_data["gene"] == gene) & (fibro_data["individual"] == ind)]
        row[f"{ind} TPM (mean)"] = sub["TPM"].mean() if len(sub) > 0 else 0
        row[f"{ind} TPM (SD)"] = sub["TPM"].std() if len(sub) > 0 else 0
    # Group means
    wt_inds = [ind for ind in INDIV_ORDER if INDIV_GENO[ind] == "WT"]
    aff_inds = [ind for ind in INDIV_ORDER if INDIV_GENO[ind] == "Affected"]
    wt_means = [row[f"{ind} TPM (mean)"] for ind in wt_inds]
    aff_means = [row[f"{ind} TPM (mean)"] for ind in aff_inds]
    row["WT mean TPM"] = np.mean(wt_means)
    row["Affected mean TPM"] = np.mean(aff_means)
    if row["WT mean TPM"] > 0:
        row["Fold change (Aff/WT)"] = row["Affected mean TPM"] / row["WT mean TPM"]
    else:
        row["Fold change (Aff/WT)"] = np.nan
    linked_rows_fibro.append(row)

linked_fibro_df = pd.DataFrame(linked_rows_fibro)

# Long-read gene-level TPM from IsoSeq pigeon v49 pipeline (whole-transcriptome)
isoseq_gene = pd.read_excel(
    "./data/pacbio-rna/pigeon_v49_rerun/results/isoseq_gene_level_FL_TPM.xlsx",
    sheet_name="Gene-level FL_TPM")

# Column mapping: BS1=Son(WT)=III-13, BS2=Mother(WT)=II-8.5, BS3=Son(aff)=III-12, BS4=Father(aff)=II-8
lr_col_map = {
    "II-8.5": "BS2 Mother(WT) FL_TPM",
    "II-8":   "BS4 Father(aff) FL_TPM",
    "III-13": "BS1 Son(WT) FL_TPM",
    "III-12": "BS3 Son(aff) FL_TPM",
}

linked_rows_lr = []
for gene in LINKED_GENES:
    row = {"Gene": gene}
    match = isoseq_gene[isoseq_gene["Gene"] == gene]
    for ind in INDIV_ORDER:
        col = lr_col_map[ind]
        row[f"{ind} FL_TPM"] = match[col].values[0] if len(match) > 0 else 0
    row["FL count"] = int(match["Total FL"].values[0]) if len(match) > 0 else 0
    wt_vals = [row[f"{ind} FL_TPM"] for ind in INDIV_ORDER if INDIV_GENO[ind] == "WT"]
    aff_vals = [row[f"{ind} FL_TPM"] for ind in INDIV_ORDER if INDIV_GENO[ind] == "Affected"]
    row["WT mean FL_TPM"] = np.mean(wt_vals)
    row["Affected mean FL_TPM"] = np.mean(aff_vals)
    if row["WT mean FL_TPM"] > 0:
        row["Fold change (Aff/WT)"] = row["Affected mean FL_TPM"] / row["WT mean FL_TPM"]
    else:
        row["Fold change (Aff/WT)"] = np.nan
    linked_rows_lr.append(row)

linked_lr_df = pd.DataFrame(linked_rows_lr)

# ══════════════════════════════════════════════════
# TABLE 4 (optional): Same for iPSCs
# ══════════════════════════════════════════════════
ipsc_files = glob.glob(os.path.join(IPSC_ROOT, "*/quant.sf"))
ipsc_data = load_gene_tpm(ipsc_files, tx_map, lambda f: os.path.basename(os.path.dirname(f)))
ipsc_data["genotype"] = ipsc_data["sample"].apply(lambda s: "Hetero" if s.startswith("hetero") else "WT")

# Exclude outlier samples
lrrfip1_tpm = ipsc_data[ipsc_data["gene"] == "LRRFIP1"].set_index("sample")["TPM"]
bad_samples = lrrfip1_tpm[lrrfip1_tpm < 10].index.tolist()
ipsc_data = ipsc_data[~ipsc_data["sample"].isin(bad_samples)]

# Clone-level means then group means
ipsc_data["clone"] = ipsc_data["sample"].apply(lambda s: s.split("_Repl")[0])

linked_rows_ipsc = []
for gene in LINKED_GENES:
    row = {"Gene": gene}
    gene_data = ipsc_data[ipsc_data["gene"] == gene]
    # Clone means
    clone_means = gene_data.groupby(["clone", "genotype"])["TPM"].mean().reset_index()
    wt_clones = clone_means[clone_means["genotype"] == "WT"]["TPM"]
    het_clones = clone_means[clone_means["genotype"] == "Hetero"]["TPM"]
    row["WT mean TPM"] = wt_clones.mean() if len(wt_clones) > 0 else 0
    row["WT SD"] = wt_clones.std() if len(wt_clones) > 0 else 0
    row["Hetero mean TPM"] = het_clones.mean() if len(het_clones) > 0 else 0
    row["Hetero SD"] = het_clones.std() if len(het_clones) > 0 else 0
    if row["WT mean TPM"] > 0:
        row["Fold change (Het/WT)"] = row["Hetero mean TPM"] / row["WT mean TPM"]
    else:
        row["Fold change (Het/WT)"] = np.nan
    linked_rows_ipsc.append(row)

linked_ipsc_df = pd.DataFrame(linked_rows_ipsc)

# ══════════════════════════════════════════════════
# TABLE: Pairwise delta (long-read FLNC)
# ══════════════════════════════════════════════════
all_lr_isos_delta = ISO_ORDER + ["Other"]
lr_totals = {ind: sum(lr_counts[iso][ind] for iso in all_lr_isos_delta) for ind in INDIV_ORDER}

delta_rows = []
for iso in all_lr_isos_delta:
    cts = lr_counts.get(iso, {ind: 0 for ind in INDIV_ORDER})
    if sum(cts.get(ind, 0) for ind in INDIV_ORDER) == 0:
        continue
    mother_pct = cts["II-8.5"] / lr_totals["II-8.5"] * 100 if lr_totals["II-8.5"] > 0 else 0
    father_pct = cts["II-8"] / lr_totals["II-8"] * 100 if lr_totals["II-8"] > 0 else 0
    son_wt_pct = cts["III-13"] / lr_totals["III-13"] * 100 if lr_totals["III-13"] > 0 else 0
    son_aff_pct = cts["III-12"] / lr_totals["III-12"] * 100 if lr_totals["III-12"] > 0 else 0
    delta_rows.append({
        "Isoform": iso,
        "Mother (WT) %": round(mother_pct, 2),
        "Father (aff) %": round(father_pct, 2),
        "Gen II Δ (pp)": round(father_pct - mother_pct, 2),
        "Son WT %": round(son_wt_pct, 2),
        "Son (aff) %": round(son_aff_pct, 2),
        "Gen III Δ (pp)": round(son_aff_pct - son_wt_pct, 2),
    })

delta_df = pd.DataFrame(delta_rows)

# ══════════════════════════════════════════════════
# TABLE: iPSC isoform fractions (Salmon, iso6 index)
# ══════════════════════════════════════════════════
ipsc_tpm_files = glob.glob(os.path.join(IPSC_ROOT, "*/quant.sf"))
ipsc_iso_rows = []
for qf in sorted(ipsc_tpm_files):
    sample = os.path.basename(os.path.dirname(qf))
    df_q = pd.read_csv(qf, sep="\t")
    sub = df_q[df_q["Name"].isin(ENST_COLS.keys())].set_index("Name")
    for enst, iso in ENST_COLS.items():
        if enst in sub.index:
            ipsc_iso_rows.append({"sample": sample, "Isoform": iso, "TPM": sub.loc[enst, "TPM"]})

ipsc_iso_data = pd.DataFrame(ipsc_iso_rows)
ipsc_iso_data["genotype"] = ipsc_iso_data["sample"].apply(lambda s: "Hetero" if s.startswith("hetero") else "WT")
ipsc_iso_data["clone"] = ipsc_iso_data["sample"].apply(lambda s: s.split("_Repl")[0])

# Total LRRFIP1 per sample for fractions + exclusion
ipsc_sample_total = ipsc_iso_data.groupby("sample")["TPM"].sum()
ipsc_iso_data["LRRFIP1_total"] = ipsc_iso_data["sample"].map(ipsc_sample_total)
ipsc_iso_data = ipsc_iso_data[ipsc_iso_data["LRRFIP1_total"] >= 10].copy()
ipsc_iso_data["fraction"] = ipsc_iso_data["TPM"] / ipsc_iso_data["LRRFIP1_total"]

clone_labels_ipsc = {"wt_1": "WT-1", "wt_2": "WT-2", "wt_3": "WT-3",
                     "hetero_1": "Het-1", "hetero_2": "Het-2", "hetero_3": "Het-3"}
wt_clones = sorted([c for c in ipsc_iso_data["clone"].unique() if c.startswith("wt")])
het_clones = sorted([c for c in ipsc_iso_data["clone"].unique() if c.startswith("hetero")])
ipsc_clone_order = wt_clones + het_clones

ipsc_table_rows = []
for clone in ipsc_clone_order:
    sub = ipsc_iso_data[ipsc_iso_data["clone"] == clone]
    geno = "WT" if clone.startswith("wt") else "Hetero"
    n_repl = sub["sample"].nunique()
    total_mean = sub.groupby("sample")["TPM"].sum().mean()
    row = {
        "Clone": clone_labels_ipsc.get(clone, clone),
        "Genotype": geno,
        "n replicates": n_repl,
        "Total LRRFIP1 TPM (mean)": total_mean,
    }
    for iso in ISO_ORDER:
        iso_sub = sub[sub["Isoform"] == iso]
        clone_frac_mean = iso_sub.groupby("sample")["fraction"].mean().mean()
        clone_frac_sd = iso_sub.groupby("sample")["fraction"].mean().std()
        clone_tpm_mean = iso_sub.groupby("sample")["TPM"].mean().mean()
        row[f"{iso} fraction (mean)"] = clone_frac_mean if not pd.isna(clone_frac_mean) else 0
        row[f"{iso} fraction (SD)"] = clone_frac_sd if not pd.isna(clone_frac_sd) else 0
        row[f"{iso} TPM (mean)"] = clone_tpm_mean if not pd.isna(clone_tpm_mean) else 0
    ipsc_table_rows.append(row)

ipsc_iso_df = pd.DataFrame(ipsc_table_rows)

# ══════════════════════════════════════════════════
# Isoform annotation table
# ══════════════════════════════════════════════════
ISO_TO_ENST = {v: k for k, v in ENST_COLS.items()}
annot_rows = []
for iso in ISO_ORDER:
    annot_rows.append({
        "Isoform": iso,
        "Ensembl_ID": ISO_TO_ENST[iso],
        "Gene": "LRRFIP1",
    })
annot_df = pd.DataFrame(annot_rows)

# ══════════════════════════════════════════════════
# Write Excel
# ══════════════════════════════════════════════════
outpath_main = os.path.join(OUTDIR, "main_figure_data_tables.xlsx")
with pd.ExcelWriter(outpath_main, engine="openpyxl") as writer:
    annot_df.to_excel(writer, sheet_name="Isoform annotation", index=False)
    sr_df.to_excel(writer, sheet_name="Short-read isoform fractions", index=False)
    lr_df.to_excel(writer, sheet_name="Long-read isoform counts", index=False)

# Format
wb = load_workbook(outpath_main)
header_font = Font(bold=True, size=9, name="Arial")
cell_font = Font(size=9, name="Arial")
header_fill = PatternFill("solid", fgColor="D9E1F2")
thin_border = Border(
    bottom=Side(style="thin", color="000000"),
)

for ws in wb.worksheets:
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = thin_border
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = cell_font
            if isinstance(cell.value, float):
                cell.number_format = "0.000" if cell.value < 1 else "0.0"
    for col_idx in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 16

wb.save(outpath_main)
print(f"Saved: {outpath_main}")

# Linked genes table
outpath_linked = os.path.join(OUTDIR, "linked_genes_TPM.xlsx")
with pd.ExcelWriter(outpath_linked, engine="openpyxl") as writer:
    linked_fibro_df.to_excel(writer, sheet_name="Fibroblast short-read", index=False)
    linked_lr_df.to_excel(writer, sheet_name="Fibroblast long-read (IsoSeq)", index=False)
    linked_ipsc_df.to_excel(writer, sheet_name="iPSC short-read (optional)", index=False)

wb2 = load_workbook(outpath_linked)
for ws in wb2.worksheets:
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = thin_border
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = cell_font
            if isinstance(cell.value, float):
                cell.number_format = "0.00"
    for col_idx in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 18

wb2.save(outpath_linked)
print(f"Saved: {outpath_linked}")

# Pairwise delta table (separate file)
outpath_delta = os.path.join(OUTDIR, "pairwise_delta_data.xlsx")
with pd.ExcelWriter(outpath_delta, engine="openpyxl") as writer:
    delta_df.to_excel(writer, sheet_name="Pairwise delta (FLNC)", index=False)

wb3 = load_workbook(outpath_delta)
for ws in wb3.worksheets:
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = thin_border
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = cell_font
            if isinstance(cell.value, float):
                cell.number_format = "0.00"
    for col_idx in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 16

wb3.save(outpath_delta)
print(f"Saved: {outpath_delta}")

# iPSC isoform data table (separate file)
outpath_ipsc = os.path.join(OUTDIR, "supp_table_iPSC_isoform_data.xlsx")
with pd.ExcelWriter(outpath_ipsc, engine="openpyxl") as writer:
    annot_df.to_excel(writer, sheet_name="Isoform annotation", index=False)
    ipsc_iso_df.to_excel(writer, sheet_name="iPSC isoform fractions", index=False)

wb4 = load_workbook(outpath_ipsc)
for ws in wb4.worksheets:
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = thin_border
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = cell_font
            if isinstance(cell.value, float):
                cell.number_format = "0.000" if cell.value < 1 else "0.0"
    for col_idx in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 16

wb4.save(outpath_ipsc)
print(f"Saved: {outpath_ipsc}")
