"""
Create supplementary Excel table with Regenie v3 per-test breakdown.
Parses all .regenie files and outputs p-values for SKAT, SKAT-O, ACAT-O, ACAT-V, SKAT-O/ACAT
and ADD burden per AAF bin.
"""
import os, re
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side, numbers
from openpyxl.utils import get_column_letter

RESULTS_DIR = "./data"
OUTDIR = "./data/salmon_quant_iPSC"

TIER_FILES = {
    "Primary (AAF ≤ 1%)":     "scoliosis_skato_v3_primary_scoliosis.regenie",
    "Ultra-rare (AAF ≤ 0.1%)":"scoliosis_skato_v3_ultrarare_scoliosis.regenie",
    "Near-private":            "scoliosis_skato_v3_singletons_scoliosis.regenie",
}

MASK_INFO = {
    "M1_cds": "CDS",
    "M2_coding": "CDS + UTR",
    "M3_regulatory": "UTR + promoter",
    "M4_intronic": "Intronic",
    "M5_genic": "CDS + UTR + intronic",
    "M6_all": "All regions",
}

GENE_ORDER = ["COL6A3", "MLPH", "PRLH", "RAB17", "LRRFIP1", "LRRFIP1_RBM44_overlap", "RBM44"]
SKIP_GENES = {"LRRFIP1_whole", "RBM44_whole"}

GENE_LABELS = {
    "COL6A3": "COL6A3",
    "MLPH": "MLPH",
    "PRLH": "PRLH",
    "RAB17": "RAB17",
    "LRRFIP1": "LRRFIP1 (unique)",
    "LRRFIP1_RBM44_overlap": "LRRFIP1/RBM44 overlap",
    "RBM44": "RBM44 (unique)",
}
MASK_ORDER = ["M1_cds", "M2_coding", "M3_regulatory", "M4_intronic", "M5_genic", "M6_all"]

SET_TESTS = ["SKAT", "SKATO", "ACATO", "ACATV", "SKATO-ACAT"]
TEST_LABELS = {
    "SKAT": "SKAT", "SKATO": "SKAT-O", "ACATO": "ACAT-O",
    "ACATV": "ACAT-V", "SKATO-ACAT": "SKAT-O/ACAT (primary)",
}

def parse_regenie(filepath):
    rows = []
    with open(filepath) as f:
        for line in f:
            if line.startswith("#") or line.startswith("CHROM"):
                continue
            parts = line.strip().split()
            if len(parts) < 12:
                continue
            raw_id, test_raw, log10p_str = parts[2], parts[7], parts[11]

            gene = mask = None
            for m in MASK_INFO:
                idx = raw_id.find("." + m)
                if idx >= 0:
                    gene = raw_id[:idx]
                    rest = raw_id[idx+1:]
                    mask = m
                    break
            if gene is None or gene in SKIP_GENES:
                continue

            if test_raw == "ADD":
                bin_part = rest[len(mask):].lstrip(".")
                test_name = f"burden_{bin_part}"
            elif test_raw.startswith("ADD-"):
                test_name = test_raw[4:]
            else:
                continue

            try:
                lp = float(log10p_str)
            except:
                lp = None
            rows.append({"gene": gene, "mask": mask, "test": test_name, "log10p": lp})
    return rows

# Parse all files
all_data = {}
for tier_label, fname in TIER_FILES.items():
    fpath = os.path.join(RESULTS_DIR, fname)
    if not os.path.isfile(fpath):
        print(f"WARNING: {fpath} not found")
        continue
    for r in parse_regenie(fpath):
        key = (r["gene"], r["mask"], r["test"], tier_label)
        all_data[key] = r["log10p"]

# ══════════════════════════════════════════════════
# Sheet 1: Set tests (SKAT, SKAT-O, ACAT-O, ACAT-V, SKAT-O/ACAT)
# ══════════════════════════════════════════════════
set_rows = []
for gene in GENE_ORDER:
    for mask in MASK_ORDER:
        row = {
            "Gene": GENE_LABELS.get(gene, gene),
            "Mask": MASK_INFO.get(mask, mask),
        }
        for test in SET_TESTS:
            for tier_label in TIER_FILES:
                key = (gene, mask, test, tier_label)
                lp = all_data.get(key)
                col_name = f"{TEST_LABELS[test]} — {tier_label}"
                if lp is not None:
                    row[col_name] = round(10**(-lp), 4) if lp < 4 else f"{10**(-lp):.2e}"
                else:
                    row[col_name] = "—"
        set_rows.append(row)

set_df = pd.DataFrame(set_rows)

# ══════════════════════════════════════════════════
# Sheet 2: ADD burden per AAF bin
# ══════════════════════════════════════════════════
# Collect all burden bin names
burden_bins = set()
for key in all_data:
    if key[2].startswith("burden_"):
        burden_bins.add(key[2])

burden_order = sorted(burden_bins)

burden_rows = []
for gene in GENE_ORDER:
    for mask in MASK_ORDER:
        row = {
            "Gene": GENE_LABELS.get(gene, gene),
            "Mask": MASK_INFO.get(mask, mask),
        }
        for bbin in burden_order:
            for tier_label in TIER_FILES:
                key = (gene, mask, bbin, tier_label)
                lp = all_data.get(key)
                bin_label = bbin.replace("burden_", "")
                col_name = f"{bin_label} — {tier_label}"
                if lp is not None:
                    row[col_name] = round(10**(-lp), 4) if lp < 4 else f"{10**(-lp):.2e}"
                else:
                    row[col_name] = "—"
        burden_rows.append(row)

burden_df = pd.DataFrame(burden_rows)

# ══════════════════════════════════════════════════
# Write Excel
# ══════════════════════════════════════════════════
outpath = os.path.join(OUTDIR, "supp_table_regenie_pertest.xlsx")
with pd.ExcelWriter(outpath, engine="openpyxl") as writer:
    set_df.to_excel(writer, sheet_name="Set tests (SKAT etc)", index=False)
    burden_df.to_excel(writer, sheet_name="ADD burden per AAF bin", index=False)

# Format
wb = load_workbook(outpath)
header_font = Font(bold=True, size=8, name="Arial")
cell_font = Font(size=8, name="Arial")
header_fill = PatternFill("solid", fgColor="D9E1F2")
sig_fill = PatternFill("solid", fgColor="FFF2CC")
thin_border = Border(bottom=Side(style="thin", color="000000"))

bonf_threshold = 0.05 / 42  # ~0.00119

for ws in wb.worksheets:
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = thin_border
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = cell_font
            cell.alignment = Alignment(horizontal="center")
            # Highlight significant p-values
            if isinstance(cell.value, (int, float)) and cell.value < bonf_threshold:
                cell.fill = sig_fill
                cell.font = Font(bold=True, size=8, name="Arial", color="CC0000")
    for col_idx in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 14
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 18

wb.save(outpath)
print(f"Saved: {outpath}")
print(f"Set tests sheet: {len(set_df)} rows, {len(set_df.columns)} columns")
print(f"Burden sheet: {len(burden_df)} rows, {len(burden_df.columns)} columns")
