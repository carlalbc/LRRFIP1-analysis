"""
iPSC isoform data supplementary table
"""
import pandas as pd
import numpy as np
import os, glob
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

IPSC_ROOT = "./data/salmon_quant_iPSC"
OUTDIR = os.path.join(IPSC_ROOT, "Final_results_30.06.26/figures_and_tables/Isoform_analyses/supplementary/tables")

ISO_ORDER = ["Isoform 1", "Isoform 2", "Isoform 3", "Isoform 4", "Isoform 5", "Isoform 6"]
ENST_COLS = {
    "ENST00000308482.14": "Isoform 1",
    "ENST00000863709.1":  "Isoform 2",
    "ENST00000392000.4":  "Isoform 3",
    "ENST00000244815.9":  "Isoform 4",
    "ENST00000289175.10": "Isoform 5",
    "ENST00000478958.5":  "Isoform 6",
}
ISO_TO_ENST = {v: k for k, v in ENST_COLS.items()}

# Load iPSC data
ipsc_iso_rows = []
for qf in sorted(glob.glob(os.path.join(IPSC_ROOT, "*/quant.sf"))):
    sample = os.path.basename(os.path.dirname(qf))
    df_q = pd.read_csv(qf, sep="\t")
    sub = df_q[df_q["Name"].isin(ENST_COLS.keys())].set_index("Name")
    for enst, iso in ENST_COLS.items():
        if enst in sub.index:
            ipsc_iso_rows.append({"sample": sample, "Isoform": iso, "TPM": sub.loc[enst, "TPM"]})

ipsc_iso_data = pd.DataFrame(ipsc_iso_rows)
ipsc_iso_data["genotype"] = ipsc_iso_data["sample"].apply(lambda s: "Hetero" if s.startswith("hetero") else "WT")
ipsc_iso_data["clone"] = ipsc_iso_data["sample"].apply(lambda s: s.split("_Repl")[0])

# Exclude replicate 2 samples (batch QC)
n_before = ipsc_iso_data["sample"].nunique()
ipsc_iso_data = ipsc_iso_data[~ipsc_iso_data["sample"].str.contains("Repl2")].copy()
n_after = ipsc_iso_data["sample"].nunique()
print(f"Excluded {n_before - n_after} replicate 2 samples ({n_before} -> {n_after})")

# Compute fractions
ipsc_sample_total = ipsc_iso_data.groupby("sample")["TPM"].sum()
ipsc_iso_data["LRRFIP1_total"] = ipsc_iso_data["sample"].map(ipsc_sample_total)
ipsc_iso_data["fraction"] = ipsc_iso_data["TPM"] / ipsc_iso_data["LRRFIP1_total"]

clone_labels_ipsc = {"wt_1": "WT-1", "wt_2": "WT-2", "wt_3": "WT-3",
                     "hetero_1": "Het-1", "hetero_2": "Het-2", "hetero_3": "Het-3"}
wt_clones = sorted([c for c in ipsc_iso_data["clone"].unique() if c.startswith("wt")])
het_clones = sorted([c for c in ipsc_iso_data["clone"].unique() if c.startswith("hetero")])
ipsc_clone_order = wt_clones + het_clones

ipsc_table_rows = []
for clone in ipsc_clone_order:
    sub = ipsc_iso_data[ipsc_iso_data["clone"] == clone]
    geno = "WT" if clone.startswith("wt") else "Hetero"
    n_repl = sub["sample"].nunique()
    total_mean = sub.groupby("sample")["TPM"].sum().mean()
    row = {
        "Clone": clone_labels_ipsc.get(clone, clone),
        "Genotype": geno,
        "n replicates": n_repl,
        "Total LRRFIP1 TPM (mean)": total_mean,
    }
    for iso in ISO_ORDER:
        iso_sub = sub[sub["Isoform"] == iso]
        clone_frac_mean = iso_sub.groupby("sample")["fraction"].mean().mean()
        clone_frac_sd = iso_sub.groupby("sample")["fraction"].mean().std()
        clone_tpm_mean = iso_sub.groupby("sample")["TPM"].mean().mean()
        row[f"{iso} fraction (mean)"] = clone_frac_mean if not pd.isna(clone_frac_mean) else 0
        row[f"{iso} fraction (SD)"] = clone_frac_sd if not pd.isna(clone_frac_sd) else 0
        row[f"{iso} TPM (mean)"] = clone_tpm_mean if not pd.isna(clone_tpm_mean) else 0
    ipsc_table_rows.append(row)

ipsc_iso_df = pd.DataFrame(ipsc_table_rows)

# Annotation table
annot_rows = [{"Isoform": iso, "Ensembl_ID": ISO_TO_ENST[iso], "Gene": "LRRFIP1"} for iso in ISO_ORDER]
annot_df = pd.DataFrame(annot_rows)

# Write Excel
outpath = os.path.join(OUTDIR, "supp_table_iPSC_isoform_data.xlsx")
with pd.ExcelWriter(outpath, engine="openpyxl") as writer:
    annot_df.to_excel(writer, sheet_name="Isoform annotation", index=False)
    ipsc_iso_df.to_excel(writer, sheet_name="iPSC isoform fractions", index=False)

# Format
wb = load_workbook(outpath)
header_font = Font(bold=True, size=9, name="Arial")
cell_font = Font(size=9, name="Arial")
header_fill = PatternFill("solid", fgColor="D9E1F2")
thin_border = Border(bottom=Side(style="thin", color="000000"))

for ws in wb.worksheets:
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = thin_border
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = cell_font
            if isinstance(cell.value, float):
                cell.number_format = "0.000" if cell.value < 1 else "0.0"
    for col_idx in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 16

wb.save(outpath)
print(f"Saved: {outpath}")
